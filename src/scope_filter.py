"""
Scope filters from the test team's filter workbooks (ADR 0041-0045).

A workbook row ("keep rows of companies 137/9088/501 created after X", as
Postgres SQL text or as SiteLink prose) is parsed ONCE into a ScopeFilter:
  - conditions on the table itself   (created_at >= '2024-08-01')
  - a join path of hops to a parent  (facility_id -> facilities.id)
  - conditions on each parent        (company_id IN (137, 9088, 501))

The hops ARE the workbook's JOIN conditions (JOIN parent a ON prev.fk = a.pk).
Two renderings of the same spec (ADR 0046):
  - render_joins(): real JOIN clauses + WHERE, for SQL we build ourselves
    (Silver emitter, both sides), where every base column is qualified.
  - render(): WHERE-only, `base.fk IN (SELECT .. FROM p1 JOIN p2 .. WHERE ..)`
    -- the same JOIN chain, inside one subquery -- for Bronze, whose data
    SELECTs are AI-written with unqualified columns (an outer JOIN would make
    id/created_at ambiguous).
The caller supplies how a side names tables/columns and which "active row"
predicate every joined table needs, so each side selects the same rows.

Prose rows can be interpreted by AI (from_ai_sql): the AI only rewrites the
prose as one SELECT..JOIN..WHERE statement, which goes through the same
parser and the same human review as a sheet cell. Nothing here guesses:
text that doesn't parse lands in ScopeFilter.errors and blocks generation.
"""
from __future__ import annotations

import re
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Callable, Dict, List, Optional

_COND_RE = re.compile(
    r"^\s*(?:(\w+)\.)?(\w+)\s*(>=|<=|<>|!=|=|>|<|\bnot\s+in\b|\bin\b)\s*(.+?)\s*$", re.I | re.S
)
_FROM_RE = re.compile(r"\bfrom\s+(?:\w+\.)*(\w+)(?:\s+(?:as\s+)?(?!join\b|where\b|left\b|inner\b)(\w+))?", re.I)
_JOIN_RE = re.compile(
    r"\bjoin\s+(?:\w+\.)*(\w+)(?:\s+(?:as\s+)?(?!on\b)(\w+))?\s+on\s+(\w+)\.(\w+)\s*=\s*(\w+)\.(\w+)", re.I
)
_PROSE_JOIN_TABLE_RE = re.compile(r"\bjoin\s+on\s+(\w+)\s+table\s+with\s+(\w+)", re.I)
_PROSE_JOIN_ON_RE = re.compile(r"\bjoin\s+(\w+)\s+on\s+(\w+)\.(\w+)\s*=\s*(\w+)\.(\w+)", re.I)
_PROSE_SCOPE_RE = re.compile(r"(\w+)\s+filter\b", re.I)


@dataclass
class Condition:
    column: str
    op: str      # =, >=, IN, ...
    value: str   # SQL literal text as written: '2024-08-01', (137, 9088, 501)

    def sql(self, column_sql: str) -> str:
        return f"{column_sql} {self.op} {self.value}"


@dataclass
class Hop:
    fk: str       # column on the previous table
    parent: str   # parent table name as the workbook spells it
    pk: str       # column on the parent the fk points at
    conditions: List[Condition] = field(default_factory=list)
    alias: str = ""  # as written in the sheet (f, t); rendered as p<n> when blank


@dataclass
class ScopeFilter:
    table: str
    mode: str = "scoped"            # scoped | full
    conditions: List[Condition] = field(default_factory=list)
    hops: List[Hop] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)   # blocking until fixed
    notes: List[str] = field(default_factory=list)    # shown, not blocking
    raw_text: str = ""
    tier: str = ""
    provenance: str = ""
    scope_values: List[str] = field(default_factory=list)  # workbook Codes sheet, context for AI

    @property
    def blocked(self) -> bool:
        return self.mode == "scoped" and bool(self.errors)

    def to_dict(self) -> dict:
        return asdict(self)


# ── parsing ─────────────────────────────────────────────────────────────────

def parse_condition(text: str) -> tuple:
    """'f.company_id in (1,2)' -> ('f', Condition). Raises ValueError."""
    m = _COND_RE.match(text or "")
    if not m:
        raise ValueError(f"not a <column> <operator> <value> condition: {text.strip()!r}")
    qualifier, column, op, value = m.groups()
    value = value.strip().rstrip(";").strip()
    if value.count("'") % 2:
        raise ValueError(f"unbalanced quote in {text.strip()!r}")
    if value.count("(") != value.count(")"):
        raise ValueError(f"unbalanced parentheses in {text.strip()!r}")
    op = " ".join(op.upper().split())
    if op in ("IN", "NOT IN") and not value.startswith("("):
        raise ValueError(f"IN needs a (...) list: {text.strip()!r}")
    return qualifier, Condition(column, op, value)


def _split_and(where: str) -> List[str]:
    return [p for p in re.split(r"\s+and\s+", where.strip().rstrip(";"), flags=re.I) if p.strip()]


def _parse_sql(spec: ScopeFilter, text: str) -> None:
    from_m = _FROM_RE.search(text)
    base_table, base_alias = from_m.group(1), from_m.group(2) or from_m.group(1)
    if base_table.lower() != spec.table.lower():
        spec.errors.append(f"filter text reads FROM {base_table}, not from {spec.table} (copy-paste?)")
    aliases = {base_alias.lower(): None}  # alias -> hop index (None = base table)
    tail = base_alias.lower()
    for jm in _JOIN_RE.finditer(text):
        parent, given_alias, la, lc, ra, rc = jm.groups()
        alias = (given_alias or parent).lower()
        if ra.lower() == alias and la.lower() == tail:
            fk, pk = lc, rc
        elif la.lower() == alias and ra.lower() == tail:
            fk, pk = rc, lc
        else:
            spec.errors.append(f"join to {parent} does not continue the path from {tail}: {jm.group(0)!r}")
            continue
        spec.hops.append(Hop(fk=fk, parent=parent, pk=pk, alias=(given_alias or "").lower()))
        aliases[alias] = len(spec.hops) - 1
        tail = alias
    where_m = re.search(r"\bwhere\b(.*)$", text, re.I | re.S)
    for part in _split_and(where_m.group(1) if where_m else ""):
        try:
            qualifier, cond = parse_condition(part)
        except ValueError as exc:
            spec.errors.append(str(exc))
            continue
        if qualifier is None:
            if spec.hops:
                spec.notes.append(f"unqualified column {cond.column} assumed to be on {spec.table}")
            spec.conditions.append(cond)
        elif qualifier.lower() not in aliases:
            spec.errors.append(f"unknown table alias {qualifier!r} in {part.strip()!r}")
        elif aliases[qualifier.lower()] is None:
            spec.conditions.append(cond)
        else:
            spec.hops[aliases[qualifier.lower()]].conditions.append(cond)


def _literal_list(values: List[str]) -> str:
    return "(" + ", ".join(v if re.fullmatch(r"-?\d+(\.\d+)?", v) else "'" + v.replace("'", "''") + "'"
                           for v in values) + ")"


def _parse_prose(spec: ScopeFilter, text: str, scope_values: List[str]) -> None:
    """SiteLink-style prose: 'join on sites table with site_id then join owners
    on sites.ownerid=owners.ownerid -- then scorpcode filter can be used'."""
    found = []
    for m in _PROSE_JOIN_TABLE_RE.finditer(text):
        found.append((m.start(), Hop(fk=m.group(2), parent=m.group(1), pk=m.group(2))))
    for m in _PROSE_JOIN_ON_RE.finditer(text):
        parent, la, lc, ra, rc = m.groups()
        fk, pk = (lc, rc) if ra.lower() == parent.lower() else (rc, lc)
        found.append((m.start(), Hop(fk=fk, parent=parent, pk=pk)))
    spec.hops = [h for _, h in sorted(found, key=lambda x: x[0])]

    scope_m = _PROSE_SCOPE_RE.search(text)
    words = text.split()
    scope_col = scope_m.group(1) if scope_m else (words[0] if len(words) == 1 else "")
    if not scope_col:
        spec.errors.append(f"could not read a filter from {text.strip()!r} -- enter it manually")
        return
    if not scope_values:
        spec.errors.append(f"no values for {scope_col} (workbook has no Codes sheet) -- enter them manually")
        return
    cond = Condition(scope_col, "IN", _literal_list(scope_values))
    (spec.hops[-1].conditions if spec.hops else spec.conditions).append(cond)


def parse_filter_text(table: str, text: Optional[str], scope_values: Optional[List[str]] = None) -> ScopeFilter:
    """One workbook cell -> ScopeFilter. Blank / 'compare full' -> mode full."""
    spec = ScopeFilter(table=table, raw_text=(text or "").strip())
    body = spec.raw_text.replace("--", " ")  # SiteLink prose continues after '--'
    if not body or re.search(r"\bcompare\s+full\b|^full$", body, re.I):
        spec.mode = "full"
        return spec
    if re.search(r"\bselect\b.+\bfrom\b", body, re.I | re.S):
        _parse_sql(spec, body)
    else:
        _parse_prose(spec, body, scope_values or [])
    if not spec.conditions and not spec.hops and not spec.errors:
        spec.errors.append("nothing to filter on was found -- enter it manually or mark full")
    return spec


# ── workbooks ───────────────────────────────────────────────────────────────

def _header_index(headers: List[str], *needles: str) -> Optional[int]:
    for needle in needles:
        for i, h in enumerate(headers):
            if needle in h:
                return i
    return None


def _scope_values_from_codes_sheet(wb) -> List[str]:
    """SiteLink workbook: a 'Codes' sheet with a 'Corp Code' column."""
    for ws in wb.worksheets:
        if "code" not in ws.title.lower():
            continue
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        for r_i, row in enumerate(rows):
            for c_i, cell in enumerate(row):
                if isinstance(cell, str) and "corp" in cell.lower():
                    return [str(r[c_i]).strip() for r in rows[r_i + 1:] if c_i < len(r) and r[c_i]]
    return []


def load_workbook(path) -> Dict[str, ScopeFilter]:
    """Filter workbook -> {table_name_lower: ScopeFilter}. Reads the first sheet
    that has a table column and a filter column (Postgres: 'Target Table Name'
    + 'Filter to apply.'; SiteLink: 'table' + 'column')."""
    import openpyxl  # already a dependency (excel_batch_loader)

    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    scope_values = _scope_values_from_codes_sheet(wb)
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        headers = [str(h or "").strip().lower() for h in rows[0]]
        t_i = _header_index(headers, "target table", "table")
        f_i = _header_index(headers, "filter", "condition", "column")
        if t_i is None or f_i is None or t_i == f_i:
            continue
        tier_i = _header_index(headers, "tier")
        specs: Dict[str, ScopeFilter] = {}
        for n, row in enumerate(rows[1:], start=2):
            table = str(row[t_i] or "").strip() if t_i < len(row) else ""
            if not table:
                continue
            spec = parse_filter_text(table, row[f_i] if f_i < len(row) else None, scope_values)
            spec.tier = str(row[tier_i] or "").strip() if tier_i is not None and tier_i < len(row) else ""
            spec.provenance = f"{path.name} / {ws.title} row {n}"
            spec.scope_values = list(scope_values)
            specs[table.lower()] = spec
        return specs
    raise ValueError(f"{path.name}: no sheet with a table column and a filter column")


def lookup(specs: Dict[str, ScopeFilter], *names: str) -> Optional[ScopeFilter]:
    """First spec matching any of *names* case-insensitively. Silver callers
    pass the Bronze table and the INT_-stripped Silver table."""
    for name in names:
        if name and name.lower() in specs:
            return specs[name.lower()]
    return None


# ── editing (UI) ────────────────────────────────────────────────────────────

def from_edit(original: ScopeFilter, mode: str, conditions_text: str, hop_rows: List[dict]) -> ScopeFilter:
    """Rebuild a spec from the UI editor. hop_rows: [{parent, alias, fk, pk, conditions}]
    (JOIN parent alias ON prev.fk = alias.pk) where conditions is
    'a = 1 AND b IN (...)'. Parse errors of the original text are kept only
    while the content is still exactly what was parsed."""
    spec = ScopeFilter(table=original.table, mode=mode, raw_text=original.raw_text,
                       tier=original.tier, provenance=original.provenance, notes=list(original.notes),
                       scope_values=list(original.scope_values))
    for part in _split_and(conditions_text or ""):
        try:
            spec.conditions.append(parse_condition(part)[1])
        except ValueError as exc:
            spec.errors.append(str(exc))
    for row in hop_rows:
        fk, parent, pk = (str(row.get(k) or "").strip() for k in ("fk", "parent", "pk"))
        if not (fk or parent or pk):
            continue
        if not (fk and parent and pk):
            spec.errors.append(f"join step to {parent or '?'} needs from-column, parent table and parent key")
            continue
        alias = str(row.get("alias") or "").strip().lower()
        if alias and not re.fullmatch(r"[a-z_]\w*", alias):
            spec.errors.append(f"alias {alias!r} must be a plain name")
        hop = Hop(fk=fk, parent=parent, pk=pk, alias=alias)
        for part in _split_and(str(row.get("conditions") or "")):
            try:
                hop.conditions.append(parse_condition(part)[1])
            except ValueError as exc:
                spec.errors.append(str(exc))
        spec.hops.append(hop)
    aliases = [_alias(spec, i) for i in range(len(spec.hops))]
    if len(set(aliases)) != len(aliases) or spec.table.lower() in aliases:
        spec.errors.append(f"join aliases must be unique and differ from {spec.table}: {', '.join(aliases)}")
    if (spec.conditions, spec.hops, spec.mode) == (original.conditions, original.hops, original.mode):
        spec.errors = list(original.errors) + spec.errors
    if spec.mode == "scoped" and not spec.conditions and not spec.hops:
        spec.errors.append("scoped filter has no conditions -- add one or choose full")
    return spec


def conditions_text(conditions: List[Condition]) -> str:
    return " AND ".join(c.sql(c.column) for c in conditions)


# ── AI interpretation (prose rows) ──────────────────────────────────────────

def from_ai_sql(original: ScopeFilter, sql: str, model: str = "") -> ScopeFilter:
    """Spec from the AI's rewrite of a prose row. The AI output is treated like
    a sheet cell: parsed by the same parser, errors block, human reviews it."""
    spec = parse_filter_text(original.table, sql, original.scope_values)
    spec.raw_text, spec.tier, spec.provenance = original.raw_text, original.tier, original.provenance
    spec.scope_values = list(original.scope_values)
    spec.notes.insert(0, f"interpreted by AI{f' ({model})' if model else ''} as: {sql.strip()} -- review before use")
    return spec


def to_sql(spec: ScopeFilter) -> str:
    """The spec as one sheet-style statement (what a reviewer / the AI reads)."""
    joins, where = render_joins(spec, lambda p: p, base_ref=spec.table)
    return f"SELECT COUNT(*) FROM {spec.table}{joins}" + (f"\nWHERE {where}" if where else "")


# ── rendering ───────────────────────────────────────────────────────────────

def _check(spec: Optional[ScopeFilter]) -> bool:
    if spec is None or spec.mode != "scoped":
        return False
    if spec.errors:
        raise ValueError(f"{spec.table}: scope filter has unresolved errors: {'; '.join(spec.errors)}")
    return True


def _alias(spec: ScopeFilter, i: int) -> str:
    return spec.hops[i].alias or f"p{i + 1}"


def _join_clause(spec, i, prev_ref, parent_ref, column) -> str:
    """JOIN <parent> <alias> ON <prev>.<fk> = <alias>.<pk> for hop i."""
    hop, alias = spec.hops[i], _alias(spec, i)
    prev_table = spec.hops[i - 1].parent if i else None
    return (f"JOIN {parent_ref(hop.parent)} {alias} ON {prev_ref}.{column(prev_table, hop.fk)} "
            f"= {alias}.{column(hop.parent, hop.pk)}")


def _hop_where(spec, column, parent_active) -> List[str]:
    parts: List[str] = []
    for i, hop in enumerate(spec.hops):
        alias = _alias(spec, i)
        if parent_active:
            parts.append(f"{alias}.{parent_active}")
        parts += [c.sql(f"{alias}.{column(hop.parent, c.column)}") for c in hop.conditions]
    return parts


def render_joins(
    spec: Optional[ScopeFilter],
    parent_ref: Callable[[str], str],
    column: Callable[[Optional[str], str], str] = lambda table, col: col,
    parent_active: str = "",
    base_ref: str = "",
) -> tuple:
    """('JOIN .. ON ..' clauses, WHERE predicate) for a query whose base table
    is referenced as *base_ref*; ('', '') when full/None.

    parent_ref(parent)       -> qualified parent table for this side
    column(table|None, col)  -> column name on this side (None = base table);
                                may raise KeyError when the side lacks it
    parent_active            -> predicate every joined table must also satisfy,
                                e.g. '_FIVETRAN_ACTIVE = TRUE'
    """
    if not _check(spec):
        return "", ""
    prefix = f"{base_ref}." if base_ref else ""
    joins = [_join_clause(spec, i, base_ref if i == 0 else _alias(spec, i - 1), parent_ref, column)
             for i in range(len(spec.hops))]
    where = [c.sql(prefix + column(None, c.column)) for c in spec.conditions]
    where += _hop_where(spec, column, parent_active)
    return "".join(f"\n{j}" for j in joins), " AND ".join(where)


def render(
    spec: Optional[ScopeFilter],
    parent_ref: Callable[[str], str],
    column: Callable[[Optional[str], str], str] = lambda table, col: col,
    parent_active: str = "",
    base_prefix: str = "",
) -> str:
    """WHERE-only form (no WHERE keyword); '' when full/None. The join chain
    goes inside one subquery: base.fk IN (SELECT p1.pk FROM p1 JOIN p2 .. WHERE ..)."""
    if not _check(spec):
        return ""
    parts = [c.sql(base_prefix + column(None, c.column)) for c in spec.conditions]
    if spec.hops:
        first, a1 = spec.hops[0], _alias(spec, 0)
        sub = f"SELECT {a1}.{column(first.parent, first.pk)} FROM {parent_ref(first.parent)} {a1}"
        sub += "".join(" " + _join_clause(spec, i, _alias(spec, i - 1), parent_ref, column)
                       for i in range(1, len(spec.hops)))
        hop_where = _hop_where(spec, column, parent_active)
        if hop_where:
            sub += " WHERE " + " AND ".join(hop_where)
        parts.append(f"{base_prefix}{column(None, first.fk)} IN ({sub})")
    return " AND ".join(parts)


def match_column(name: str, live_columns: List[str]) -> Optional[str]:
    """Live column matching *name* ignoring case and underscores
    (site_id ~ SiteID ~ SITE_ID). None when zero or several match."""
    key = name.replace("_", "").lower()
    hits = [c for c in live_columns if c.replace("_", "").lower() == key]
    return hits[0] if len(hits) == 1 else None


def silver_parent_column(parent: str, column: str) -> str:
    """Silver column for a parent-table column. Coalesce renames the Bronze
    primary key ID to <SINGULAR>_ID (INT_FACILITIES: ID -> FACILITY_ID, see
    docs/metadata_1.txt); every other column keeps its name.
    ponytail: naive singular (ies->y, sses->ss, s->''); the UI checks the result
    against live INT_ columns and lets a human override it."""
    if column.lower() != "id":
        return column.upper()
    p = parent.upper()
    singular = p[:-3] + "Y" if p.endswith("IES") else p[:-2] if p.endswith("SSES") else p[:-1] if p.endswith("S") else p
    return f"{singular}_ID"


def silver_base_column(mappings, bronze_column: str) -> str:
    """Silver column that passes *bronze_column* through unchanged, from the
    plan's own Coalesce mappings ("T"."ID" -> DISCOUNT_LINE_ID). KeyError when
    the Silver table has no passthrough of that column (ADR 0043 gate)."""
    want = bronze_column.replace('"', "").upper()
    for m in mappings:
        src = (m.source_column or "").strip()
        if re.fullmatch(r'"[^"]+"\."[^"]+"', src) and src.split(".")[-1].strip('"').upper() == want:
            return m.target_column
    raise KeyError(bronze_column)


def _resolver(default: Callable[[Optional[str], str], str], live: Optional[Dict[str, List[str]]],
              missing: List[str]) -> Callable[[Optional[str], str], str]:
    """Wrap a side's default naming with a live-column check. live is
    {'' (base table) | parent_lower: [columns]}; tables not in it are unchecked."""
    live = {k.lower(): v for k, v in (live or {}).items()}

    def resolve(table: Optional[str], col: str) -> str:
        try:
            name = default(table, col)
        except KeyError:
            missing.append(f"{table or 'base table'}: no column for {col}")
            return col.upper()
        cols = live.get((table or "").lower())
        if cols is None:
            return name
        hit = match_column(name, cols)
        if hit is None:
            missing.append(f"{table or 'base table'}.{name} not found")
            return name
        return hit
    return resolve


def render_bronze(spec, source_schema: str, sf_database: str, sf_schema: str,
                  target_column_for: Optional[Dict[str, str]] = None,
                  live_source: Optional[Dict[str, List[str]]] = None,
                  live_target: Optional[Dict[str, List[str]]] = None) -> tuple:
    """Source-DB -> Bronze Snowflake. Returns (source_pred, target_pred, missing).
    target_column_for: {source_col_lower: snowflake_col} from the column mapping."""
    missing: List[str] = []
    target_column_for = {k.lower(): v for k, v in (target_column_for or {}).items()}
    sf_prefix = ".".join(p for p in (sf_database, sf_schema) if p)
    src = render(spec, parent_ref=lambda p: f"{source_schema}.{p}" if source_schema else p,
                 column=_resolver(lambda t, c: c, live_source, missing))
    tgt = render(spec, parent_ref=lambda p: f"{sf_prefix}.{p.upper()}",
                 column=_resolver(lambda t, c: target_column_for.get(c.lower(), c.upper()) if t is None else c.upper(),
                                  live_target, missing),
                 parent_active="_FIVETRAN_ACTIVE = TRUE")
    return src, tgt, missing


def render_silver(spec, plan, live_bronze: Optional[Dict[str, List[str]]] = None,
                  live_silver: Optional[Dict[str, List[str]]] = None) -> dict:
    """Bronze Snowflake -> Silver (INT_<table>), as real JOINs (ADR 0046).
    Returns {bronze_joins, bronze_where, silver_joins, silver_where, missing}.
    Joined tables: Bronze <PARENT> + _FIVETRAN_ACTIVE, Silver INT_<PARENT> + IS_CURRENT."""
    from silver.silver_sql_emitter import bronze_aliases

    missing: List[str] = []
    q = lambda name: f'"{name}"'
    bronze_col = _resolver(lambda t, c: c.upper(), live_bronze, missing)
    silver_col = _resolver(
        lambda t, c: silver_base_column(plan.mappings, c) if t is None else silver_parent_column(t, c),
        live_silver, missing,
    )
    bronze_joins, bronze_where = render_joins(
        spec, parent_ref=lambda p: f'"{plan.source_database}"."{plan.source_schema}"."{p.upper()}"',
        column=lambda t, c: q(bronze_col(t, c)), parent_active='"_FIVETRAN_ACTIVE" = TRUE',
        base_ref=f'"{bronze_aliases(plan)[0]}"')
    silver_joins, silver_where = render_joins(
        spec, parent_ref=lambda p: f'"{plan.target_database}"."{plan.target_schema}"."INT_{p.upper()}"',
        column=lambda t, c: q(silver_col(t, c)), parent_active='"IS_CURRENT" = TRUE',
        base_ref=f'"{plan.target_table}"')
    return dict(bronze_joins=bronze_joins, bronze_where=bronze_where,
                silver_joins=silver_joins, silver_where=silver_where, missing=missing)


if __name__ == "__main__":
    s = parse_filter_text(
        "phone_numbers",
        "select count(*) from phone_numbers j join tenants t on j.owner_id=t.id "
        "join facilities f on t.facility_id=f.id where j.owner_type in ('Tenant') and f.company_id in (137,9088,501);",
    )
    assert not s.errors and [h.parent for h in s.hops] == ["tenants", "facilities"]
    assert render(s, lambda p: f"public.{p}") == (
        "owner_type IN ('Tenant') AND owner_id IN (SELECT t.id FROM public.tenants t "
        "JOIN public.facilities f ON t.facility_id = f.id WHERE f.company_id IN (137,9088,501))"
    )
    assert render_joins(s, lambda p: p, base_ref="j") == (
        "\nJOIN tenants t ON j.owner_id = t.id\nJOIN facilities f ON t.facility_id = f.id",
        "j.owner_type IN ('Tenant') AND f.company_id IN (137,9088,501)",
    )
    print("ok")
